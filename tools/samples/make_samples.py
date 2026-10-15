#!/usr/bin/env python3
"""Generate the ElixirOffice sample documents shipped with the example app, the
admin panel preview and the WOPI dummy file (replacing the upstream's samples).

    python3 tools/samples/make_samples.py [outdir]     (default: ./samples)

Writes sample.docx (short newsletter), sample.xlsx (small household budget with a
chart), sample.pptx (3 slides) and sample.pdf (the newsletter). All content is
invented and international (English, no country-specific data). Requires
python-docx, python-pptx, openpyxl, reportlab and Pillow (pip install them in a venv).
Output is committed in the document-templates repository (sample/); tests must
never write into those files (copy to a temporary directory first).
"""
import datetime as dt
import os
import sys

from docx import Document
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt as PPt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = sys.argv[1] if len(sys.argv) > 1 else "samples"
WHEN = dt.datetime(2026, 1, 15, 9, 0, 0)
ACCENT = RGBColor(0x1F, 0x6F, 0xEB)

TITLE = "Welcome to ElixirOffice"
LEAD = ("ElixirOffice is a private, AI-ready office suite that runs in your browser and "
        "feels like the tools you already know. This short newsletter is a sample document: "
        "edit it, comment on it, or start a new one from the Home page.")
POINTS = [
    "Write letters, reports and newsletters with the usual shortcuts.",
    "Build budgets and charts in the spreadsheet editor.",
    "Present your ideas with clean, consistent slides.",
    "Keep your files on your own server: nothing leaves without your say.",
]
EVENTS = [("12 Feb", "Open workshop: first steps", "Room A"),
          ("26 Feb", "Tips and shortcuts", "Online"),
          ("11 Mar", "Share and review together", "Room B")]
CLOSING = "Questions or ideas? Reply to this newsletter and we will gladly help."


def make_docx(path):
    d = Document()
    cp = d.core_properties
    cp.author, cp.title, cp.created, cp.modified = "ElixirOffice", TITLE, WHEN, WHEN
    t = d.add_paragraph(TITLE, style="Title")
    t.runs[0].font.color.rgb = ACCENT
    d.add_paragraph("Sample newsletter - January 2026", style="Subtitle")
    d.add_paragraph(LEAD)
    d.add_heading("What you can do", level=1)
    for p in POINTS:
        d.add_paragraph(p, style="List Bullet")
    d.add_heading("Upcoming events", level=1)
    tb = d.add_table(rows=1, cols=3, style="Light Grid Accent 1")
    for c, h in zip(tb.rows[0].cells, ("Date", "Event", "Place")):
        c.text = h
    for row in EVENTS:
        cells = tb.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    d.add_paragraph()
    d.add_paragraph(CLOSING)
    d.save(path)


def make_xlsx(path):
    wb = Workbook()
    wb.properties.creator, wb.properties.title = "ElixirOffice", "Household budget (sample)"
    wb.properties.created = wb.properties.modified = WHEN
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Household budget - first half of the year (sample data)"
    ws["A1"].font = Font(bold=True, size=14)
    head = ["Month", "Housing", "Food", "Transport", "Leisure", "Total"]
    ws.append([])
    ws.append(head)
    data = [("Jan", 900, 320, 85, 120), ("Feb", 900, 295, 90, 80), ("Mar", 900, 340, 80, 150),
            ("Apr", 900, 310, 95, 110), ("May", 900, 330, 88, 170), ("Jun", 900, 305, 92, 140)]
    for i, r in enumerate(data, start=4):
        ws.append(list(r) + [f"=SUM(B{i}:E{i})"])
    last = 3 + len(data)
    ws.append(["Total"] + [f"=SUM({c}4:{c}{last})" for c in "BCDEF"])
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F6FEB")
        c.alignment = Alignment(horizontal="center")
    for c in ws[last + 1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, max_row=last + 1, min_col=2, max_col=6):
        for c in row:
            c.number_format = "#,##0"
    ws.column_dimensions["A"].width = 12
    for c in "BCDEF":
        ws.column_dimensions[c].width = 12
    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title = "Monthly spending by category"
    ch.add_data(Reference(ws, min_col=2, max_col=5, min_row=3, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=last))
    ch.y_axis.delete = ch.x_axis.delete = False
    ch.x_axis.axPos, ch.y_axis.axPos = "b", "l"  # the editor draws a column chart wrongly with both axes on the left
    ch.height, ch.width = 8.5, 17
    ws.add_chart(ch, "H3")
    wb.save(path)


def make_pptx(path):
    p = Presentation()
    p.core_properties.author, p.core_properties.title = "ElixirOffice", TITLE
    p.core_properties.created = p.core_properties.modified = WHEN
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = TITLE
    s.placeholders[1].text = "A three-slide sample presentation"
    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "What you can do"
    tf = s.placeholders[1].text_frame
    tf.text = POINTS[0]
    for t in POINTS[1:]:
        tf.add_paragraph().text = t
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Workshop attendance (sample data)"
    cd = CategoryChartData()
    cd.categories = ["Jan", "Feb", "Mar", "Apr"]
    cd.add_series("Attendees", (24, 31, 28, 40))
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.9), Inches(1.8), Inches(8.2), Inches(5), cd)
    gf.chart.has_legend = True
    gf.chart.legend.position, gf.chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    p.save(path)


def make_pdf(path):
    c = canvas.Canvas(path, pagesize=A4, invariant=1)
    c.setTitle(TITLE)
    c.setAuthor("ElixirOffice")
    w, h = A4
    y = h - 90
    c.setFillColorRGB(0.12, 0.44, 0.92)
    c.setFont("Helvetica-Bold", 26)
    c.drawString(60, y, TITLE)
    c.setFillColorRGB(0.35, 0.35, 0.35)
    c.setFont("Helvetica-Oblique", 12)
    c.drawString(60, y - 24, "Sample newsletter - January 2026")
    c.setFillColorRGB(0, 0, 0)
    c.setFont("Helvetica", 11)
    y -= 64
    words, line = LEAD.split(), ""
    for wd in words:
        if c.stringWidth(line + " " + wd, "Helvetica", 11) > w - 120:
            c.drawString(60, y, line.strip()); y -= 16; line = ""
        line += " " + wd
    c.drawString(60, y, line.strip())
    y -= 34
    c.setFont("Helvetica-Bold", 15)
    c.drawString(60, y, "What you can do")
    c.setFont("Helvetica", 11)
    for t in POINTS:
        y -= 20
        c.drawString(70, y, "-  " + t)
    y -= 38
    c.setFont("Helvetica-Bold", 15)
    c.drawString(60, y, "Upcoming events")
    c.setFont("Helvetica", 11)
    for r in EVENTS:
        y -= 20
        c.drawString(70, y, r[0]); c.drawString(140, y, r[1]); c.drawString(400, y, r[2])
    y -= 44
    c.drawString(60, y, CLOSING)
    c.save()


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    for ext, fn in (("docx", make_docx), ("xlsx", make_xlsx), ("pptx", make_pptx), ("pdf", make_pdf)):
        fn(os.path.join(OUT, "sample." + ext))
        print("wrote", os.path.join(OUT, "sample." + ext))
