#!/usr/bin/env python3
# SPDX-License-Identifier: AGPL-3.0-only
"""Generate the deterministic .docx documents used by the performance harness.

    python3 tools/perf/gen_docs.py [outdir]     (default: tools/perf/docs)

perf-10p.docx   10 pages of text (headings, body paragraphs, a bullet list per page)
perf-50p.docx   50 pages: text, a 12x5 table every 5 pages, an image every 7 pages
perf-100p.docx 100 pages: text, a 15x6 table every 3 pages, a 2000x1300 photo-like JPEG every 4 pages

Each page ends with an explicit page break and holds less than a page of content, so the page
count does not depend on font metrics. Requires python-docx and Pillow
(python3 -m pip install --user python-docx).
"""
import io
import os
import random
import sys

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFilter

WORDS = ("le la les un une des et en dans pour avec sur par au aux du de que qui est sont "
         "classe élève exercice leçon fraction nombre problème calcul texte lecture écriture "
         "projet séance objectif compétence évaluation méthode exemple document tableau image "
         "chapitre partie résultat question réponse correction consigne activité groupe travail "
         "mesure longueur surface volume durée vitesse histoire géographie science nature "
         "observer comparer expliquer décrire justifier calculer rédiger lire présenter organiser").split()


def sentence(rng, n=None):
    n = n or rng.randint(8, 18)
    w = [rng.choice(WORDS) for _ in range(n)]
    return " ".join(w).capitalize() + rng.choice([".", ".", ".", " ?", " !"])


def paragraph(rng, sentences):
    return " ".join(sentence(rng) for _ in range(sentences))


def image(rng, w, h, quality=82, blur=3):
    """Photo-like JPEG (noise + blurred shapes) so that it does not compress to nothing."""
    img = Image.effect_noise((w, h), rng.randint(30, 70)).convert("RGB")
    d = ImageDraw.Draw(img)
    for _ in range(25):
        x, y = rng.randint(0, w), rng.randint(0, h)
        r = rng.randint(w // 20, w // 4)
        d.ellipse((x - r, y - r, x + r, y + r), fill=tuple(rng.randint(0, 255) for _ in range(3)))
    img = img.filter(ImageFilter.GaussianBlur(blur))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=quality)
    buf.seek(0)
    return buf


def table(doc, rng, rows, cols):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    for c in range(cols):
        t.cell(0, c).text = rng.choice(WORDS).capitalize()
    for r in range(1, rows):
        for c in range(cols):
            t.cell(r, c).text = str(rng.randint(1, 999)) if c else " ".join(rng.choice(WORDS) for _ in range(3))
    return t


def build(path, pages, table_every=0, image_every=0, seed=1, photo=(1400, 900, 82, 3)):
    rng = random.Random(seed)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for p in range(1, pages + 1):
        doc.add_heading(f"Partie {p} : {sentence(rng, 4)[:-1].lower()}", level=1 if p % 5 == 1 else 2)
        has_table = table_every and p % table_every == 0
        has_image = image_every and p % image_every == 0
        body = 1 if (has_table and has_image) else (2 if (has_table or has_image) else 4)
        for _ in range(body):
            doc.add_paragraph(paragraph(rng, rng.randint(3, 5)))
        if not has_table and not has_image:
            for _ in range(3):
                doc.add_paragraph(sentence(rng), style="List Bullet")
        if has_table:
            table(doc, rng, 12 if pages < 100 else 15, 5 if pages < 100 else 6)
        if has_image:
            doc.add_picture(image(rng, *photo), width=Cm(11 if has_table else 14))
        if p < pages:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)
    return os.path.getsize(path)


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs")
    os.makedirs(out, exist_ok=True)
    for name, args in [("perf-10p.docx", dict(pages=10, seed=10)),
                       ("perf-50p.docx", dict(pages=50, table_every=5, image_every=7, seed=50)),
                       ("perf-100p.docx", dict(pages=100, table_every=3, image_every=4, seed=100, photo=(2000, 1300, 88, 1)))]:
        size = build(os.path.join(out, name), **args)
        print(f"{name}: {size / 1024:.0f} KiB")


if __name__ == "__main__":
    main()
